"""Écriture de la transcription dans différents formats : .txt, .md, .docx."""

from __future__ import annotations

from pathlib import Path


def _fmt_duration(seconds: float | None) -> str:
    if not seconds:
        return "inconnue"
    seconds = int(seconds)
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h} h {m:02d} min {s:02d} s"
    if m:
        return f"{m} min {s:02d} s"
    return f"{s} s"


def _header_lines(meta: dict) -> list[str]:
    return [
        f"Langue détectée : {meta.get('language', 'inconnue')}",
        f"Durée de l'audio : {_fmt_duration(meta.get('duration'))}",
        f"Fichier source : {meta.get('source', 'inconnu')}",
    ]


def export_txt(path: str | Path, text: str, meta: dict) -> None:
    Path(path).write_text(text.strip() + "\n", encoding="utf-8")


def export_md(path: str | Path, text: str, meta: dict) -> None:
    lines = ["# Transcription", ""]
    lines += [f"- {line}" for line in _header_lines(meta)]
    lines += ["", "---", "", text.strip(), ""]
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def export_docx(path: str | Path, text: str, meta: dict) -> None:
    from docx import Document  # import différé : dépendance lourde

    doc = Document()
    doc.add_heading("Transcription", level=0)
    for line in _header_lines(meta):
        doc.add_paragraph(line, style="Intense Quote" if False else None)
    doc.add_paragraph("")
    for paragraph in text.strip().split("\n"):
        doc.add_paragraph(paragraph)
    doc.save(str(path))


# Correspondance extension -> fonction d'export.
EXPORTERS = {
    ".txt": export_txt,
    ".md": export_md,
    ".docx": export_docx,
}


def export(path: str | Path, text: str, meta: dict) -> None:
    """Choisit l'exporteur d'après l'extension du fichier (défaut : .txt)."""
    suffix = Path(path).suffix.lower()
    EXPORTERS.get(suffix, export_txt)(path, text, meta)
